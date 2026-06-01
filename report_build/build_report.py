# -*- coding: utf-8 -*-
"""
Builds the professional WeatherBot Software-Engineering project report (.docx).
Auto TOC, List of Figures and List of Tables use Word field codes + the
Caption style, so they populate when the document fields are updated in Word.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(__file__)
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "..", "WeatherBot_Project_Report.docx")
URL = "https://khayamtariq.github.io/Weatherbot/"

NAVY  = RGBColor(0x0B, 0x1F, 0x3A)
BLUE  = RGBColor(0x1D, 0x4E, 0xD8)
GREY  = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_BG = "1D4ED8"
HDR_BG = "0B1F3A"

doc = Document()

# ---------------- base styles ------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for i, sz, col in [(1, 16, NAVY), (2, 13, BLUE), (3, 12, NAVY)]:
    st = doc.styles[f"Heading {i}"]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.color.rgb = col
    st.font.bold = True
    st.paragraph_format.space_before = Pt(12 if i == 1 else 8)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def shade_cell_text(cell, color=WHITE, bold=True):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = color
            r.font.bold = bold


def field(paragraph, instr, placeholder):
    r = paragraph.add_run()
    fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end")
    r._r.append(fb); r._r.append(it); r._r.append(fs); r._r.append(t); r._r.append(fe)


def body(text, italic=False, align="justify", bold=False, color=None, size=None,
         space_after=6):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.italic = italic; r.bold = bold
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def h1(text): doc.add_heading(text, level=1)
def h2(text): doc.add_heading(text, level=2)
def h3(text): doc.add_heading(text, level=3)


def caption(label, text):
    """label = 'Figure' or 'Table'; uses SEQ so lists auto-build."""
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"{label} ")
    field(p, f" SEQ {label} \\* ARABIC ", "0")
    rest = p.add_run(f". {text}")
    for r in p.runs:
        r.font.size = Pt(9.5); r.font.italic = True; r.font.color.rgb = GREY


def figure(img, cap, width=6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(os.path.join(ASSETS, img), width=Inches(width))
    caption("Figure", cap)


def table(headers, rows, caption_text, widths=None, header_bg=HDR_BG, font_size=10):
    caption("Table", caption_text)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = htext
        set_cell_bg(hdr[i], header_bg)
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(font_size); r.font.bold = True; r.font.color.rgb = WHITE
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)
            if ri % 2 == 1:
                set_cell_bg(cells[i], "EEF2F7")
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "1D4ED8")
    pbdr.append(bottom); pPr.append(pbdr)


# =======================================================================
# COVER PAGE
# =======================================================================
def cover():
    for _ in range(2):
        doc.add_paragraph()
    body("SOFTWARE ENGINEERING PROJECT REPORT", align="center", bold=True,
         color=BLUE, size=13, space_after=2)
    hr()
    for _ in range(2):
        doc.add_paragraph()
    body("WeatherBot", align="center", bold=True, color=NAVY, size=40, space_after=2)
    body("An AI-Style, Multi-Modal Conversational Weather Assistant",
         align="center", italic=True, color=GREY, size=15, space_after=4)
    doc.add_paragraph()

    # logo-ish block
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("🌤  WeatherBot"); r.font.size = Pt(22); r.bold = True; r.font.color.rgb = BLUE

    for _ in range(2):
        doc.add_paragraph()
    body("Technical Documentation & System Design", align="center", bold=True,
         color=NAVY, size=14, space_after=2)
    body("Final-Year Software Engineering Project", align="center", italic=True,
         color=GREY, size=12)
    for _ in range(3):
        doc.add_paragraph()

    # info table
    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    info = [("Project Title", "WeatherBot — Conversational Weather Chatbot"),
            ("Live Deployment", URL),
            ("Document Type", "Software Engineering Project Report"),
            ("Document Version", "2.0 (Revised & Expanded)")]
    for i, (k, v) in enumerate(info):
        c0, c1 = t.rows[i].cells
        c0.text = k; c1.text = v
        set_cell_bg(c0, HDR_BG)
        for p in c0.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = WHITE; r.font.size = Pt(10.5)
        for p in c1.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5)
    for row in t.rows:
        row.cells[0].width = Inches(2.2); row.cells[1].width = Inches(4.0)
    doc.add_page_break()


# =======================================================================
# FRONT MATTER
# =======================================================================
def certificate():
    h1("Certificate")
    body("This is to certify that the project entitled “WeatherBot — An "
         "AI-Style, Multi-Modal Conversational Weather Assistant” is a record "
         "of bona fide work carried out as part of the Software Engineering "
         "curriculum. The system described in this report has been designed, "
         "implemented and deployed as a fully functional, publicly accessible web "
         "application, and the documentation faithfully reflects the actual "
         "implementation of the delivered software.")
    body("The work embodied in this report is original and has been completed in "
         "accordance with the academic requirements for the submission of a "
         "final-year Software Engineering project. The live system is available "
         "for verification at the deployment URL listed below.")
    doc.add_paragraph()
    body(f"Live System: {URL}", bold=True, color=BLUE, align="left")
    doc.add_paragraph(); doc.add_paragraph()
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells
    a.text = "\n\n_________________________\nProject Supervisor"
    b.text = "\n\n_________________________\nHead of Department"
    for c in (a, b):
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def acknowledgement():
    h1("Acknowledgement")
    body("We would like to express our sincere gratitude to everyone who "
         "supported the development of WeatherBot. We are especially thankful to "
         "our project supervisor for continuous guidance, constructive feedback "
         "and encouragement throughout every phase of the software development "
         "life cycle.")
    body("We gratefully acknowledge the open-source and open-data communities "
         "whose freely available services made this project possible — in "
         "particular the Open-Meteo project for weather and geocoding data, "
         "BigDataCloud for client-side reverse geocoding, the OpenStreetMap "
         "contributors for global map tiles, and the Leaflet library for "
         "interactive mapping. Their commitment to free, key-less APIs is the "
         "foundation of WeatherBot’s zero-configuration philosophy.")
    body("Finally, we thank our faculty members, peers and families for their "
         "patience, motivation and valuable suggestions, which contributed "
         "significantly to the successful completion of this work.")
    doc.add_page_break()


def abstract():
    h1("Abstract")
    body("WeatherBot is a lightweight, multi-modal conversational weather "
         "assistant that runs entirely inside a standard web browser. Presented "
         "as a modern, frosted-glass chat interface, it allows users to obtain "
         "live weather information for any location on Earth using four distinct "
         "interaction modes: natural-language typing, voice (speech-to-text and "
         "text-to-speech), GPS-based geolocation, and direct selection on an "
         "interactive world map. Each of these inputs is reduced to a uniform "
         "internal representation — a location (latitude, longitude and "
         "place name) together with a time intent (today, tomorrow, or a seven-"
         "day outlook) — and routed through a single weather-reporting "
         "routine that retrieves and renders the result.")
    body("Architecturally, WeatherBot is a serverless, single-file web "
         "application. It contains no back-end, no database, no build step and no "
         "API keys; all orchestration is performed on the client in vanilla "
         "JavaScript, while live data is fetched directly from free, public, "
         "key-less web services over HTTPS. A compact rule-based natural-language "
         "parser extracts user intent without relying on a machine-learning "
         "model, keeping the application instantaneous, transparent and offline-"
         "capable in its logic. The system is publicly deployed and operates "
         f"continuously at {URL}.")
    body("This report documents the complete engineering of WeatherBot using a "
         "formal software-engineering methodology. It presents the problem "
         "statement, objectives and scope; a full requirements specification; "
         "the system architecture and design expressed through standard UML and "
         "data-flow diagrams; detailed implementation and component descriptions; "
         "a structured testing and validation plan; and a critical discussion of "
         "results, advantages, limitations and future enhancements.")
    doc.add_page_break()


def toc_pages():
    h1("Table of Contents")
    p = doc.add_paragraph()
    field(p, ' TOC \\o "1-3" \\h \\z \\u ',
          'Right-click here and choose “Update Field” to generate the Table of Contents.')
    doc.add_page_break()

    h1("List of Figures")
    p = doc.add_paragraph()
    field(p, ' TOC \\h \\z \\c "Figure" ',
          'Right-click here and choose “Update Field” to generate the List of Figures.')
    doc.add_page_break()

    h1("List of Tables")
    p = doc.add_paragraph()
    field(p, ' TOC \\h \\z \\c "Table" ',
          'Right-click here and choose “Update Field” to generate the List of Tables.')
    doc.add_page_break()


# =======================================================================
# CHAPTERS
# =======================================================================
def ch_introduction():
    h1("1. Introduction")
    h2("1.1 Background")
    body("Access to timely and reliable weather information is part of everyday "
         "decision-making — from choosing what to wear and whether to carry "
         "an umbrella, to planning travel and outdoor activities. While numerous "
         "weather websites and mobile applications exist, many of them are heavy, "
         "advertisement-laden, require installation or accounts, and present "
         "information through dense dashboards rather than a natural, "
         "conversational exchange. There is value in demonstrating that a "
         "polished, modern and genuinely useful weather experience can be "
         "delivered with nothing more than standard browser technologies and a "
         "handful of free public services.")
    body("WeatherBot answers this need. It reframes weather retrieval as a chat "
         "conversation: the user simply asks a question — by typing, "
         "speaking, sharing their location, or tapping a map — and the "
         "assistant replies with a friendly, formatted, optionally spoken answer. "
         "The application is intentionally engineered to be portable, transparent "
         "and effortless to deploy.")
    h2("1.2 Purpose of the Document")
    body("This document is the formal Software Engineering project report for "
         "WeatherBot. Its purpose is to capture, in a structured and verifiable "
         "manner, the complete engineering of the system: the problem it solves, "
         "its requirements, its analysis and design, its implementation, the "
         "testing performed against it, and a critical evaluation of the outcome. "
         "The report is intended for academic evaluators, future maintainers and "
         "any reader wishing to understand how the deployed system works.")
    h2("1.3 Document Conventions")
    bullet(" denote validated functional or design facts drawn directly from the "
           "source implementation.", bold_lead="Code-derived statements")
    bullet(" are numbered sequentially and referenced in the body text; the List "
           "of Figures is auto-generated.", bold_lead="Figures and Tables")
    bullet(" such as report(), parse() and geocode() refer to the actual "
           "JavaScript functions in the application.", bold_lead="Monospace identifiers")


def ch_problem():
    h1("2. Problem Statement")
    body("Conventional weather applications frequently present several recurring "
         "shortcomings that motivate this project:")
    bullet(" Many solutions require account creation, API keys, app installation "
           "or paid tiers before a user can obtain a simple forecast.",
           bold_lead="Friction and configuration. ")
    bullet(" Information is often delivered through cluttered dashboards rather "
           "than a natural question-and-answer interaction.",
           bold_lead="Non-conversational interfaces. ")
    bullet(" Few lightweight tools allow the user to choose how they provide a "
           "location — by name, by voice, by GPS, or by pointing at a map "
           "— within a single consistent interface.",
           bold_lead="Limited input flexibility. ")
    bullet(" Many web apps require a server, a database and a build pipeline, "
           "increasing the cost and complexity of hosting.",
           bold_lead="Heavy deployment footprint. ")
    body("The problem this project addresses can therefore be stated as follows: "
         "How can a single, self-contained web application deliver accurate, "
         "live, multi-modal weather information through a natural conversational "
         "interface, without any back-end infrastructure, credentials or "
         "installation, while remaining responsive, accessible and continuously "
         "available to anyone with a web browser?")


def ch_objectives():
    h1("3. Objectives")
    body("The objectives of the WeatherBot project are:")
    numbered("To design and implement a conversational weather assistant that "
             "interprets plain-English questions and replies in a friendly, "
             "human-readable format.")
    numbered("To support four interchangeable input modes — typed text, "
             "voice, GPS geolocation and interactive map selection — that "
             "all converge on a single weather-reporting pipeline.")
    numbered("To provide three time horizons for every location: current "
             "conditions, a next-day outlook, and a seven-day forecast.")
    numbered("To retrieve live meteorological and geospatial data exclusively "
             "from free, public, key-less APIs over secure HTTPS connections.")
    numbered("To build the entire product as a single, serverless HTML file that "
             "requires no database, no build step and no configuration.")
    numbered("To deliver a responsive, accessible and visually polished user "
             "interface that adapts to both desktop and mobile devices.")
    numbered("To handle network, permission and input errors gracefully, always "
             "responding with a helpful message rather than failing silently.")
    numbered("To deploy the system publicly so that it is continuously available "
             f"24/7 at {URL}.")


def ch_scope():
    h1("4. Scope of the System")
    h2("4.1 In-Scope")
    bullet("Conversational retrieval of current weather, next-day outlook and "
           "seven-day forecasts for any geocodable location worldwide.")
    bullet("Four input modes: natural-language text, voice input/output, GPS "
           "geolocation, and interactive map selection.")
    bullet("Client-side natural-language parsing for place name and time intent.")
    bullet("Mapping of numeric weather codes to human-readable descriptions and "
           "emoji icons.")
    bullet("Graceful error handling and a responsive, accessible interface.")
    bullet("Public, zero-configuration deployment on a static host.")
    h2("4.2 Out-of-Scope")
    bullet("User accounts, authentication, personalisation or persistent "
           "history (the application is intentionally stateless).")
    bullet("Server-side processing, databases or any self-hosted back-end.")
    bullet("Hourly forecasts, severe-weather push alerts and historical climate "
           "archives (identified as future enhancements).")
    bullet("Languages other than English for parsing and speech.")
    bullet("Native mobile (Android/iOS) applications.")


def ch_literature():
    h1("5. Literature Review")
    body("The design of WeatherBot draws on several established areas of software "
         "engineering and web technology.")
    h2("5.1 Conversational and Multi-Modal Interfaces")
    body("Conversational user interfaces reduce the cognitive load of "
         "structured forms by allowing users to express intent in natural "
         "language. Multi-modal interaction — combining text, voice and "
         "direct manipulation — is widely recognised to improve "
         "accessibility and user satisfaction, since users can choose the "
         "modality most convenient for their context. WeatherBot applies these "
         "principles by funnelling four heterogeneous input modes into one "
         "uniform internal contract.")
    h2("5.2 Rule-Based versus Machine-Learning NLU")
    body("Natural-language understanding (NLU) can be implemented with heavy "
         "machine-learning language models or with lightweight rule-based "
         "parsers. For a narrowly scoped domain such as “which place and "
         "which day”, a transparent rule-based approach offers instantaneous "
         "performance, no model download, no inference cost and fully "
         "deterministic behaviour. WeatherBot deliberately adopts this approach, "
         "trading broad linguistic coverage for speed, simplicity and offline "
         "logic.")
    h2("5.3 Serverless, API-Composition Architectures")
    body("A growing class of applications are built as thin clients that compose "
         "several independent public web services rather than maintaining their "
         "own back-end. This “JAMstack/serverless” style improves "
         "portability and lowers operational cost. WeatherBot follows this model: "
         "the browser itself acts as client, orchestrator and presentation layer, "
         "calling Open-Meteo, BigDataCloud and OpenStreetMap directly.")
    h2("5.4 Progressive Enhancement and Graceful Degradation")
    body("Modern browser APIs such as Geolocation and the Web Speech API are not "
         "uniformly supported. Best practice dictates progressive enhancement: "
         "core functionality must work everywhere, while advanced features "
         "degrade gracefully with a clear notice where unsupported. WeatherBot "
         "implements this for both voice and location features.")


def ch_analysis():
    h1("6. System Analysis")
    h2("6.1 Existing Systems")
    body("Typical existing weather solutions fall into three groups: (a) full "
         "web portals that are feature-rich but heavy and advertisement-driven; "
         "(b) native mobile apps that require installation and permissions; and "
         "(c) raw weather APIs that demand keys and developer effort. None "
         "combine a conversational, multi-modal, zero-configuration experience in "
         "a single shareable file.")
    h2("6.2 Proposed System")
    body("The proposed system, WeatherBot, is a single-file browser application "
         "that unifies four input modes behind a chat interface and composes free "
         "public APIs to deliver live forecasts. It eliminates installation, "
         "accounts, keys and servers, and is deployable to any static host or "
         "even opened directly from the file system.")
    h2("6.3 Feasibility Analysis")
    bullet(" The system relies only on mature, standardised browser APIs "
           "(Fetch, Geolocation, Web Speech, DOM) and established free services, "
           "all of which are proven and available.", bold_lead="Technical. ")
    bullet(" Because there is no server, database or paid API, hosting cost is "
           "effectively zero; the deployed instance runs on free static hosting.",
           bold_lead="Economic. ")
    bullet(" The conversational interface lowers the learning curve; users "
           "interact in plain language with immediate feedback.",
           bold_lead="Operational. ")
    bullet(" The single-file, dependency-light design can be built and "
           "maintained quickly and is trivial to deploy and reason about.",
           bold_lead="Schedule. ")
    h2("6.4 Functional Requirements")
    body("The functional requirements define what the system must do. They are "
         "summarised in Table 1.")
    table(
        ["ID", "Requirement", "Description"],
        [
            ["FR-1", "Typed weather query", "The system shall accept a free-text question and return the relevant weather information."],
            ["FR-2", "Voice query", "The system shall capture spoken input via the Web Speech API and process it as a text query."],
            ["FR-3", "Spoken replies", "The system shall optionally read bot replies aloud using text-to-speech."],
            ["FR-4", "GPS location", "The system shall obtain the user's coordinates via the Geolocation API and report local weather."],
            ["FR-5", "Map selection", "The system shall let the user pick any point on an interactive map and report its weather."],
            ["FR-6", "Time intents", "The system shall provide current, next-day and 7-day forecasts."],
            ["FR-7", "Geocoding", "The system shall convert a place name to coordinates and coordinates to a place name."],
            ["FR-8", "Code translation", "The system shall map numeric weather codes to descriptions and emoji icons."],
            ["FR-9", "Suggestion chips", "The system shall present quick-reply chips to guide common questions."],
            ["FR-10", "Error feedback", "The system shall respond to every failure with a clear, helpful chat message."],
        ],
        "Functional requirements of WeatherBot.",
        widths=[0.7, 1.9, 3.6],
    )
    h2("6.5 Non-Functional Requirements")
    body("The non-functional requirements define quality attributes and "
         "constraints. They are summarised in Table 2.")
    table(
        ["ID", "Attribute", "Requirement"],
        [
            ["NFR-1", "Performance", "Local parsing and rendering shall be instantaneous; only geocoding and forecast calls incur network latency."],
            ["NFR-2", "Portability", "The application shall run as a single HTML file on any modern browser or static host."],
            ["NFR-3", "Usability", "The interface shall be conversational, responsive and adapt to phone and desktop screens."],
            ["NFR-4", "Reliability", "The system shall degrade gracefully on network, permission or unsupported-feature failures."],
            ["NFR-5", "Security & Privacy", "All external calls shall use HTTPS; no keys, accounts or persistent user data shall be used."],
            ["NFR-6", "Maintainability", "Logic shall be organised into cohesive function groups around a single reporting hub."],
            ["NFR-7", "Availability", "The deployed system shall be continuously accessible 24/7 from a public URL."],
            ["NFR-8", "Accessibility", "The system shall offer voice interaction and clear visual hierarchy for inclusive use."],
        ],
        "Non-functional requirements of WeatherBot.",
        widths=[0.7, 1.6, 3.9],
    )


def ch_techstack():
    h1("7. Technology Stack")
    body("WeatherBot deliberately uses a minimal, dependency-light stack. "
         "Table 3 lists each technology and the role it plays in the system.")
    table(
        ["Technology", "Category", "Role in WeatherBot"],
        [
            ["HTML5", "Structure", "Defines the chat panel, header, message area, map overlay, input row and tool buttons."],
            ["CSS3", "Presentation", "Implements the responsive glassmorphism theme, animated background blobs and message styling."],
            ["JavaScript (ES6+)", "Application logic", "Implements parsing, geocoding, forecast retrieval, rendering and all orchestration on the client."],
            ["Fetch API", "Networking", "Performs asynchronous HTTPS requests to the external services."],
            ["Geolocation API", "Browser capability", "Provides the user's GPS coordinates on request."],
            ["Web Speech API", "Browser capability", "Provides speech-to-text input and text-to-speech output."],
            ["Leaflet 1.9.4", "Mapping library", "Renders the interactive world map for point selection."],
            ["Open-Meteo APIs", "External data", "Supply geocoding and live weather/forecast data, key-free."],
            ["BigDataCloud API", "External data", "Provides client-side reverse geocoding (coordinates to place name)."],
            ["OpenStreetMap", "External data", "Supplies the map tiles displayed by Leaflet."],
            ["GitHub Pages", "Deployment", "Hosts the static application publicly and continuously over HTTPS."],
        ],
        "Technology stack and the role of each component.",
        widths=[1.5, 1.3, 3.4],
    )


def ch_architecture():
    h1("8. System Architecture")
    body("Architecturally, WeatherBot is a thin client that orchestrates several "
         "independent public services. Everything inside the browser is organised "
         "into a presentation layer (HTML/CSS), an application-logic layer "
         "(JavaScript), and a set of native browser capabilities. Outside the "
         "browser sit four external services, each reached over HTTPS. Because "
         "there is no back-end of our own, the browser itself is the application "
         "server, client and orchestrator all at once.")
    h2("8.1 High-Level Architecture")
    figure("fig_architecture.png",
           "High-level layered architecture: a single HTML file orchestrates four "
           "external services directly from the browser.", width=6.2)
    body("As shown in Figure 1, the presentation layer renders the interface and "
         "captures events; the application-logic layer performs parsing, "
         "geocoding, forecast retrieval and rendering; and the native browser "
         "capabilities provide location, speech, networking and DOM access. The "
         "external services are reached only through the Fetch API over HTTPS.")
    body("This design has three important consequences. First, the application "
         "is stateless and serverless — there is nothing to deploy beyond a "
         "static file. Second, all data is fetched on demand and nothing is "
         "stored. Third, the user’s browser is the trust boundary: location "
         "and microphone access are requested directly by the browser and never "
         "pass through a server we control.")
    h2("8.2 Component Architecture")
    figure("fig_components.png",
           "Component architecture: cohesive modules revolve around the central "
           "report() reporting hub.", width=5.6)
    body("Figure 2 shows the logical grouping of the code. Although the entire "
         "application lives in one file, it is divided into cohesive components "
         "— the user interface, the search/NLU module, the weather-API "
         "module, the data-processing module, the location-services module, the "
         "voice module, the error-handling module and the responsive-design layer "
         "— all of which interact through the central report() hub.")
    h2("8.3 Architectural Rationale")
    body("Every interaction is ultimately reduced to a location object (a "
         "latitude, a longitude and a place name) plus a time intent (today, "
         "tomorrow or the week). No matter how the user starts — typing, "
         "speaking, GPS or the map — the application funnels that input into "
         "the same core report() routine. Because all four input modes converge "
         "on a single function, the system is easy to extend: adding a new way to "
         "choose a location requires nothing more than producing a location "
         "object and calling report(). The data contract between components is "
         "deliberately small and uniform.")


def ch_design():
    h1("9. System Design")
    body("This chapter expresses the design of WeatherBot through standard "
         "data-flow and UML diagrams. Each diagram has been recreated to ensure "
         "clean connectors, correct notation and consistency with the actual "
         "implementation.")
    h2("9.1 Context Diagram")
    figure("fig_context.png",
           "Context diagram: WeatherBot as a single process exchanging data with "
           "the user and four external services.", width=6.3)
    body("The context diagram (Figure 3) places WeatherBot at the centre as a "
         "single process. The user is the only human external entity, while "
         "Open-Meteo Geocoding, Open-Meteo Forecast, BigDataCloud and "
         "OpenStreetMap are external service entities. Each connector is labelled "
         "with the data that flows across it.")
    h2("9.2 Level-0 Data Flow Diagram")
    figure("fig_dfd0.png",
           "Level-0 DFD: the four major processes of WeatherBot and their data "
           "stores.", width=6.4)
    body("The Level-0 DFD (Figure 4) decomposes the system into four major "
         "processes — (1.0) Capture & Parse Input, (2.0) Resolve Location, "
         "(3.0) Fetch & Process Forecast, and (4.0) Render & Speak Reply — "
         "together with the in-memory weather-code lookup table (D1).")
    h2("9.3 Level-1 Data Flow Diagram")
    figure("fig_dfd1.png",
           "Level-1 DFD: decomposition of Process 1.0 (Capture & Parse Input).",
           width=6.4)
    body("The Level-1 DFD (Figure 5) decomposes Process 1.0 into its "
         "sub-processes: detecting the input mode (1.1), normalising free text "
         "(1.2), detecting the time intent (1.3) and extracting the city name "
         "(1.4), supported by the filler-word and keyword lists (D2).")
    h2("9.4 Use-Case Diagram")
    figure("fig_usecase.png",
           "Use-case diagram: the user and the principal use cases of the system.",
           width=5.4)
    body("Figure 6 captures the principal use cases available to the single "
         "actor, the User: asking by typing, asking by voice, using GPS, picking "
         "a point on the map, and hearing replies read aloud.")
    h2("9.5 Sequence Diagram")
    figure("fig_sequence.png",
           "Sequence diagram for the typed query “Will it rain in Paris "
           "tomorrow?”.", width=6.4)
    body("The sequence diagram (Figure 7) traces a typed query through the "
         "system: the UI calls parse() to obtain the city and time intent, "
         "geocode() resolves the coordinates via Open-Meteo, report() fetches the "
         "forecast, and the reply is rendered and optionally spoken. The two "
         "network calls are asynchronous; a temporary typing bubble is shown "
         "while they are in flight.")
    h2("9.6 Activity Diagram")
    figure("fig_activity.png",
           "Activity diagram: control flow and decision points of a weather "
           "request.", width=4.6)
    body("Figure 8 models the control flow as activities and decisions: the mode "
         "is determined, free text is parsed and geocoded (or coordinates are "
         "reverse-geocoded), and — if a location is found — the forecast "
         "is fetched, formatted and rendered; otherwise an error path returns the "
         "user to the input stage.")
    h2("9.7 Module (Class-Style) Diagram")
    figure("fig_class.png",
           "Module diagram: the logical classes/modules, their data and "
           "operations.", width=6.4)
    body("Although JavaScript organises the code as functions rather than "
         "classes, Figure 9 presents the logical modules in class notation: the "
         "Location data structure, the NLU parser, the GeocodingService, the "
         "WeatherService and the UIController, with their principal attributes "
         "and operations.")
    h2("9.8 Component Diagram")
    figure("fig_component_uml.png",
           "UML component diagram: deployable components and their interfaces.",
           width=6.2)
    body("Figure 10 shows the system as UML components — the Chat UI, the "
         "Weather Engine, the Browser APIs and the External Web APIs — and "
         "the interfaces through which they communicate.")


def ch_implementation():
    h1("10. Implementation Details")
    body("WeatherBot is implemented as a single self-contained HTML document "
         "comprising three concerns: structure (HTML), presentation (CSS) and "
         "behaviour (JavaScript). This section describes how each contributes to "
         "the system.")
    h2("10.1 Frontend — HTML")
    body("The HTML defines the application shell: a header with branding, a "
         "scrollable chat container, a map overlay panel, and a form containing a "
         "row of tool buttons (location, map, microphone and speaker) and a text "
         "input with a Send button. Semantic structure and ARIA-friendly titles "
         "on the buttons support accessibility, and the viewport meta tag enables "
         "responsive behaviour on mobile devices.")
    h2("10.2 Frontend — CSS")
    body("The CSS implements the signature glassmorphism theme. Four animated, "
         "colourful “blobs” drift behind the chat panel and are blurred "
         "through it using a CSS backdrop-filter, producing a frosted-glass "
         "effect. Message bubbles use distinct gradients for user and bot, a pop "
         "animation on entry, and a pulsing red state for the microphone while "
         "listening. The layout uses flexbox and relative units so it adapts "
         "fluidly between phone and desktop form factors.")
    h2("10.3 Frontend — JavaScript")
    body("The JavaScript layer contains all of the application logic, organised "
         "into cohesive groups of functions:")
    bullet(" capture raw input (text, voice transcript, GPS fix or map tap).",
           bold_lead="Event handlers ")
    bullet(" turns a free-text question into a city name and a time intent using "
           "transparent text rules.", bold_lead="parse() ")
    bullet(" convert a place name or coordinates into a concrete "
           "location object.", bold_lead="geocode() / reverseGeocode() ")
    bullet(" the single hub that takes a location plus time intent, fetches the "
           "forecast and renders the answer.", bold_lead="report() ")
    bullet(" display every message and the suggestion chips.",
           bold_lead="addMsg() / addChips() ")
    bullet(" optionally reads the reply aloud through the speech layer.",
           bold_lead="speak() ")
    h2("10.4 Backend / API Integration")
    body("WeatherBot has no back-end of its own; instead it integrates four "
         "external services directly from the browser. The integration follows a "
         "consistent request lifecycle.")
    h3("10.4.1 API Request Lifecycle")
    numbered("A location is obtained either by geocoding a place name or by "
             "reverse-geocoding coordinates.")
    numbered("report() issues a Fetch request to the Open-Meteo forecast "
             "endpoint, passing the latitude and longitude and requesting current "
             "conditions plus seven days of daily data with automatic timezone "
             "handling.")
    numbered("The asynchronous response is awaited and parsed as JSON.")
    numbered("Relevant fields are extracted (temperature, apparent temperature, "
             "humidity, wind speed, weather code and precipitation probability).")
    numbered("The numeric weather code is translated locally into a description "
             "and emoji via the lookup table, and the reply is formatted "
             "according to the time intent and rendered into the chat.")
    h3("10.4.2 JSON Response Processing")
    body("The forecast response contains a current block and a daily block. For "
         "the “today” intent the current block is used; for "
         "“tomorrow” the second element of the daily arrays is read; "
         "and for the “week” intent the daily arrays are iterated to "
         "build a seven-line forecast. Table 4 lists the external services and "
         "the data each provides.")
    table(
        ["Service", "Endpoint role", "Key-free?", "Data returned"],
        [
            ["Open-Meteo Geocoding", "Place name → coordinates", "Yes", "Latitude, longitude, name, admin region, country."],
            ["Open-Meteo Forecast", "Coordinates → weather", "Yes", "Current conditions and 7-day daily forecast with weather codes."],
            ["BigDataCloud", "Coordinates → place name", "Yes", "City, locality, principal subdivision and country."],
            ["OpenStreetMap (via Leaflet)", "Map tiles", "Yes", "Raster map tiles for the interactive map."],
        ],
        "External services and data sources used by WeatherBot.",
        widths=[1.7, 1.7, 0.8, 2.2],
    )
    body("The mapping from numeric weather codes to descriptions and emoji is "
         "implemented as an in-memory lookup table. Table 5 shows a representative "
         "sample of this mapping.")
    table(
        ["Code", "Description", "Icon", "Code", "Description", "Icon"],
        [
            ["0", "Clear sky", "☀️", "61", "Light rain", "\U0001F326️"],
            ["2", "Partly cloudy", "⛅", "65", "Heavy rain", "\U0001F327️"],
            ["3", "Overcast", "☁️", "71", "Light snow", "\U0001F328️"],
            ["45", "Foggy", "\U0001F32B️", "95", "Thunderstorm", "⛈️"],
        ],
        "Representative sample of the weather-code mapping table.",
        widths=[0.7, 1.7, 0.6, 0.7, 1.7, 0.6],
    )


def ch_workflow():
    h1("11. Working Methodology and Complete System Workflow")
    body("This chapter explains the complete operational workflow of WeatherBot, "
         "stage by stage. For each stage the input, the processing performed and "
         "the resulting output are stated explicitly. A defining characteristic "
         "of the system is that four very different input methods all feed into "
         "the same weather pipeline, as illustrated in Figure 11.")
    figure("fig_inputs.png",
           "All four input modes converge on a single weather-reporting pipeline.",
           width=6.2)

    steps = [
        ("Step 1 — Open the Website",
         "The user navigates to the deployment URL in a browser.",
         "The single HTML file is loaded; the chat interface initialises and "
         "displays a greeting message with suggestion chips.",
         "A ready, interactive chat panel."),
        ("Step 2 — Enter a City / Choose a Mode",
         "A typed question, a spoken phrase, a GPS request, or a map tap.",
         "The relevant event handler captures the raw input and identifies the "
         "input mode.",
         "Raw input passed into the pipeline (text string or coordinates)."),
        ("Step 3 — Input Validation",
         "The captured text or coordinates.",
         "For text, parse() strips filler words; if no place name remains, the "
         "bot asks the user to name a city. Coordinates are accepted directly.",
         "Either a validated query (city + intent) or a clarifying prompt."),
        ("Step 4 — Weather API Request Generation",
         "A validated city name (text modes) or coordinates (GPS/map modes).",
         "For text, geocode() builds an Open-Meteo geocoding URL; for "
         "coordinates, reverseGeocode() builds a BigDataCloud URL. report() then "
         "builds the Open-Meteo forecast URL.",
         "Fully-formed HTTPS request URLs."),
        ("Step 5 — API Communication",
         "The constructed request URLs.",
         "The Fetch API sends asynchronous HTTPS requests to the external "
         "services; a temporary typing bubble is shown while awaiting responses.",
         "Raw HTTP responses from the services."),
        ("Step 6 — Response Handling",
         "The raw HTTP responses.",
         "Responses are awaited and parsed as JSON; a missing or empty geocoding "
         "result triggers a “check the spelling” message.",
         "Parsed JSON objects (location and forecast)."),
        ("Step 7 — Data Extraction",
         "The parsed forecast JSON.",
         "Relevant fields are extracted — temperature, apparent temperature, "
         "humidity, wind speed, weather code and precipitation probability — "
         "for the requested time horizon.",
         "A structured set of weather values."),
        ("Step 8 — Code Translation & Formatting",
         "The numeric weather code and extracted values.",
         "The code is mapped to a description and emoji via the lookup table, and "
         "the reply text is composed according to the time intent.",
         "A friendly, formatted reply string."),
        ("Step 9 — UI Rendering",
         "The formatted reply string.",
         "addMsg() creates a bot bubble, appends it to the chat and auto-scrolls "
         "to the newest message; suggestion chips may be added.",
         "A rendered reply visible in the chat."),
        ("Step 10 — Optional Speech Output",
         "The rendered bot reply (when speech output is enabled).",
         "speak() cleans the text of emoji and markup, expands units into "
         "speakable words and uses text-to-speech to read it aloud.",
         "An audible spoken reply."),
        ("Step 11 — Forecast Drill-Down",
         "A tap on a “Tomorrow” or “7-day forecast” chip.",
         "report() is re-invoked for the same location with a different time "
         "intent, without re-locating the user.",
         "An updated forecast reply for the chosen horizon."),
        ("Step 12 — Error Handling",
         "Any failure during the above stages (network, permission or no match).",
         "The corresponding handler removes the typing bubble and renders a "
         "specific, helpful message; technical errors are logged to the console.",
         "A graceful, informative error message in the chat."),
    ]
    for title, inp, proc, outp in steps:
        h3(title)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run("Input: "); r.bold = True; r.font.color.rgb = BLUE
        p.add_run(inp)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        r = p.add_run("Processing: "); r.bold = True; r.font.color.rgb = BLUE
        p.add_run(proc)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        r = p.add_run("Output: "); r.bold = True; r.font.color.rgb = BLUE
        p.add_run(outp)


def ch_components():
    h1("12. Component Description")
    body("This chapter describes each component of WeatherBot. For every "
         "component its purpose, responsibilities, inputs, outputs, internal "
         "processing and interaction with other components are stated.")

    comps = [
        ("12.1 User Interface Component",
         "Render the chat experience and capture user interaction.",
         "Display message bubbles and suggestion chips, manage the map overlay, "
         "maintain tool-button states, and route events to the logic layer.",
         "User clicks, taps, text entry and form submissions.",
         "Rendered chat messages, chips and visual state changes.",
         "Builds DOM elements, applies CSS classes and animations, auto-scrolls "
         "the chat and toggles button states.",
         "Sends captured input to the Search/NLU and Location modules and "
         "receives formatted text from report() for display."),
        ("12.2 Search / NLU Module",
         "Interpret free-text questions.",
         "Detect the time intent and isolate the place name from a natural-"
         "language message.",
         "A free-text string from typed or voice input.",
         "A structured object containing the city name and the time intent.",
         "Applies keyword tests for time intent and removes filler words, "
         "punctuation and weather vocabulary to leave the place name.",
         "Feeds its output to the Geocoding portion of the Location module and "
         "ultimately to report()."),
        ("12.3 Weather API Module",
         "Retrieve live forecast data.",
         "Construct the Open-Meteo forecast request and obtain current and daily "
         "weather data.",
         "Latitude and longitude.",
         "A parsed JSON forecast object.",
         "Builds the request URL with the required current and daily parameters "
         "and automatic timezone handling, then issues an asynchronous Fetch "
         "call.",
         "Invoked by report(); its output is consumed by the Data Processing "
         "module."),
        ("12.4 Data Processing Module",
         "Transform raw forecast data into a human-readable reply.",
         "Extract relevant fields, translate weather codes and format the reply "
         "by time intent.",
         "The forecast JSON and the requested time intent.",
         "A formatted, friendly reply string.",
         "Reads the appropriate current or daily values, maps codes to text and "
         "emoji via the lookup table, rounds numeric values and assembles the "
         "message.",
         "Receives data from the Weather API module and passes output to the UI "
         "component for rendering."),
        ("12.5 Forecast Module",
         "Provide the three time horizons.",
         "Produce current, next-day and seven-day views from a single forecast "
         "response.",
         "The parsed forecast and the selected time intent.",
         "The corresponding current, tomorrow or weekly reply.",
         "Selects the current block, the second daily element, or iterates all "
         "daily elements, depending on the intent.",
         "Implemented within report(); re-invoked by drill-down chips."),
        ("12.6 Location Services Module",
         "Resolve locations from coordinates and the map.",
         "Obtain GPS coordinates, reverse-geocode coordinates to a place name, "
         "and handle interactive map selection.",
         "A GPS request or a map tap.",
         "A location object (coordinates plus place name).",
         "Calls the Geolocation API or Leaflet map events and the BigDataCloud "
         "reverse-geocoding service, with fallbacks when data is missing.",
         "Supplies location objects directly to report(), bypassing parsing and "
         "forward geocoding."),
        ("12.7 Voice / Speech Module",
         "Enable spoken interaction.",
         "Capture spoken questions (speech-to-text) and read replies aloud "
         "(text-to-speech).",
         "Microphone audio (input) or a bot reply string (output).",
         "A recognised text query, or audible speech.",
         "Uses the Web Speech API for recognition and synthesis, cleans text of "
         "emoji and expands units for natural pronunciation, and manages the "
         "listening visual state.",
         "Feeds recognised text into the same submission path as typed input; "
         "reads replies generated by report()."),
        ("12.8 Error Handling Module",
         "Ensure graceful failure.",
         "Detect and respond to network errors, permission denials, unsupported "
         "features and unmatched locations.",
         "Exceptions and error conditions from any stage.",
         "A specific, helpful chat message; console diagnostics.",
         "Wraps network calls in try/catch, inspects geolocation and speech error "
         "codes, and removes the typing indicator before showing the message.",
         "Interacts with every module, always returning control to the UI "
         "component with a user-facing message."),
        ("12.9 Responsive Design Module",
         "Adapt the interface to any device.",
         "Maintain layout, readability and usability across screen sizes.",
         "Viewport dimensions and device characteristics.",
         "A fluid, correctly proportioned interface.",
         "Uses flexible CSS layout, relative units, the viewport meta tag and a "
         "bounded maximum panel width.",
         "Underpins the UI component on all devices."),
    ]
    for title, purpose, resp, inp, outp, internal, interact in comps:
        h2(title)
        for lead, val in [("Purpose: ", purpose), ("Responsibilities: ", resp),
                          ("Inputs: ", inp), ("Outputs: ", outp),
                          ("Internal Processing: ", internal),
                          ("Interaction: ", interact)]:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(lead); r.bold = True; r.font.color.rgb = BLUE
            p.add_run(val)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def ch_ui():
    h1("13. User Interface Description")
    body("The interface is a single frosted-glass chat panel centred over an "
         "animated, multi-colour background. The layout is a vertical stack: a "
         "header, a scrollable message area, an optional map overlay, and a fixed "
         "input area with a row of tool buttons. Figure 12 presents an annotated "
         "wireframe of the interface.")
    figure("fig_wireframe.png",
           "Annotated wireframe of the WeatherBot interface.", width=4.0)
    h2("13.1 Screen: Main Chat Interface")
    body("Purpose: the primary and only screen of the application, through which "
         "all interaction occurs.")
    body("Components present: branding header; scrollable conversation area with "
         "bot and user bubbles and suggestion chips; tool-button row (location, "
         "map, microphone, speaker); text input field; and Send button.")
    body("User actions: type a question and press Send; tap a suggestion chip; "
         "press the location, map, microphone or speaker buttons.")
    body("System responses: a user bubble appears for the query, a typing "
         "indicator is shown during network calls, and a formatted bot bubble "
         "(optionally spoken) presents the result.")
    h2("13.2 Screen: Interactive Map Overlay")
    body("Purpose: allow the user to select any point on Earth to obtain its "
         "weather.")
    body("Components present: an instruction bar, a Close button, and a full "
         "interactive Leaflet/OpenStreetMap map.")
    body("User actions: open the map via the map tool button, pan and zoom, tap "
         "a location, or close the overlay.")
    body("System responses: a marker is placed, the overlay closes, the tapped "
         "point is reverse-geocoded, and its weather is reported with drill-down "
         "chips.")
    h2("13.3 User-Interface Elements")
    body("Table 6 describes every visible interface element and its behaviour.")
    table(
        ["Element", "Type", "Behaviour"],
        [
            ["Header / logo", "Static branding", "Displays the WeatherBot name and tagline."],
            ["Chat area", "Scrollable region", "Shows the conversation; auto-scrolls to the newest message."],
            ["Bot bubble", "Message", "Frosted-glass bubble aligned left for system replies."],
            ["User bubble", "Message", "Gradient bubble aligned right for user messages."],
            ["Suggestion chips", "Quick actions", "Submit a preset question or run a custom action (e.g. My location)."],
            ["Location button", "Tool", "Requests GPS coordinates and reports local weather."],
            ["Map button", "Tool", "Opens the interactive map overlay for point selection."],
            ["Microphone button", "Tool", "Starts voice input; pulses red while listening."],
            ["Speaker button", "Tool", "Toggles text-to-speech for replies (muted/active)."],
            ["Text input", "Field", "Accepts free-text questions; submits on Send or Enter."],
            ["Send button", "Action", "Submits the typed question to the pipeline."],
        ],
        "Description of the visible user-interface elements.",
        widths=[1.5, 1.3, 3.4],
    )
    h2("13.4 Notable UI Behaviours")
    bullet("New messages animate in with a subtle pop, and the chat auto-scrolls "
           "to the newest bubble.")
    bullet("User and bot bubbles are visually distinct (gradient versus frosted "
           "glass) and aligned to opposite sides.")
    bullet("The microphone button pulses red while listening; the speaker button "
           "toggles between muted and active states.")
    bullet("Suggestion chips can either submit a pre-written question or run a "
           "custom action.")


def ch_accessibility():
    h1("14. Website Accessibility and Deployment")
    body("WeatherBot is publicly deployed and continuously available. This "
         "chapter explains how users access and use the live system.")
    h2("14.1 Deployment URL")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(URL)
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = BLUE
    body("The application is hosted on GitHub Pages, a static hosting service "
         "that serves the single HTML file over HTTPS. Because the application is "
         "serverless and stateless, it runs continuously (24/7) with no "
         "maintenance, and every visit loads the latest deployed version.")
    h2("14.2 Accessing the Website Through a Browser")
    body("A user accesses WeatherBot by entering the deployment URL into the "
         "address bar of any modern web browser, or by following a shared link. "
         "No installation, sign-up, account or API key is required — the "
         "interface is ready for interaction the moment the page loads.")
    h2("14.3 Supported Browsers")
    body("WeatherBot runs on all modern, standards-compliant browsers. Core "
         "functionality (typed queries, map selection and forecasts) works "
         "across browsers, while certain advanced features depend on browser "
         "support and degrade gracefully where unavailable.")
    table(
        ["Browser", "Core chat & maps", "Voice input/output", "Notes"],
        [
            ["Google Chrome", "Fully supported", "Supported", "Recommended for full voice support."],
            ["Microsoft Edge", "Fully supported", "Supported", "Full feature support."],
            ["Mozilla Firefox", "Fully supported", "Limited", "Speech support varies by version/platform."],
            ["Safari", "Fully supported", "Limited", "Speech support varies; HTTPS required for mic."],
            ["Mobile browsers", "Fully supported", "Device-dependent", "Responsive layout; GPS commonly available."],
        ],
        "Browser compatibility and feature support.",
        widths=[1.4, 1.4, 1.3, 2.1],
    )
    h2("14.4 User Workflow: From Opening the Site to Obtaining Weather")
    numbered("Open the deployment URL in a browser; the chat interface loads "
             "with a greeting and suggestion chips.")
    numbered("Provide a location by typing a question, tapping a chip, speaking, "
             "sharing GPS, or selecting a point on the map.")
    numbered("Grant browser permission for location or microphone if prompted "
             "(only required for those specific modes).")
    numbered("Read the formatted reply in the chat, and optionally enable the "
             "speaker to hear it read aloud.")
    numbered("Drill down into the next-day or seven-day forecast using the "
             "provided chips.")
    h2("14.5 Responsiveness on Desktop and Mobile")
    body("The interface is fully responsive. On desktop it appears as a centred "
         "chat panel with a bounded maximum width; on mobile it expands to a "
         "comfortable full-height column. Flexible CSS layout and the viewport "
         "meta tag ensure that text, bubbles, buttons and the map remain legible "
         "and usable across screen sizes and orientations.")
    h2("14.6 Internet Connectivity Requirements")
    body("An active internet connection is required, because all live data "
         "(geocoding, forecasts, reverse geocoding and map tiles) is fetched on "
         "demand from external services. The application logic itself — "
         "parsing and rendering — runs locally and instantaneously; only the "
         "data retrieval depends on the network. If connectivity fails, the "
         "system reports a clear network-error message.")
    h2("14.7 User Interaction Process")
    body("Interaction is conversational and immediate: the user expresses a "
         "request through any supported mode, the system acknowledges with a user "
         "bubble and a typing indicator, and a formatted reply appears within "
         "moments. Guidance chips reduce the need to type, and every error "
         "condition is communicated in plain language, keeping the experience "
         "approachable for non-technical users.")


def ch_testing():
    h1("15. Testing and Validation")
    body("WeatherBot was validated through structured testing across functional, "
         "interface, integration (API), input-validation, error-handling and "
         "browser-compatibility dimensions. The results are presented in the "
         "following tables.")
    h2("15.1 Functional Testing")
    table(
        ["ID", "Test case", "Expected result", "Status"],
        [
            ["FT-1", "Typed query “Weather in Tokyo”", "Current conditions for Tokyo are shown.", "Pass"],
            ["FT-2", "Query with “tomorrow”", "Next-day outlook is returned.", "Pass"],
            ["FT-3", "Query with “forecast”", "Seven-day forecast is returned.", "Pass"],
            ["FT-4", "Suggestion chip tapped", "Corresponding query is submitted automatically.", "Pass"],
            ["FT-5", "Drill-down chip (7-day)", "Weekly forecast for the same location appears.", "Pass"],
        ],
        "Functional test cases and results.",
        widths=[0.6, 2.3, 2.6, 0.7],
    )
    h2("15.2 User-Interface Testing")
    table(
        ["ID", "Test case", "Expected result", "Status"],
        [
            ["UT-1", "New message added", "Bubble animates in and chat auto-scrolls.", "Pass"],
            ["UT-2", "Microphone active", "Button pulses red while listening.", "Pass"],
            ["UT-3", "Speaker toggled", "Icon switches between muted and active.", "Pass"],
            ["UT-4", "Map overlay opened", "Map displays and is pannable/zoomable.", "Pass"],
            ["UT-5", "Resize to mobile width", "Layout adapts without overflow.", "Pass"],
        ],
        "User-interface test cases and results.",
        widths=[0.6, 2.3, 2.6, 0.7],
    )
    h2("15.3 API (Integration) Testing")
    table(
        ["ID", "Test case", "Expected result", "Status"],
        [
            ["AT-1", "Geocode valid city", "Coordinates and place name returned.", "Pass"],
            ["AT-2", "Forecast request", "Current and daily data returned and parsed.", "Pass"],
            ["AT-3", "Reverse-geocode coordinates", "Readable place name returned.", "Pass"],
            ["AT-4", "Map tiles load", "OpenStreetMap tiles render in the overlay.", "Pass"],
        ],
        "API integration test cases and results.",
        widths=[0.6, 2.3, 2.6, 0.7],
    )
    h2("15.4 Input-Validation Testing")
    table(
        ["ID", "Test case", "Expected result", "Status"],
        [
            ["IV-1", "Empty input submitted", "Submission ignored; no error thrown.", "Pass"],
            ["IV-2", "Query with no place name", "Bot asks the user to name a city.", "Pass"],
            ["IV-3", "Filler-heavy phrasing", "Place name correctly isolated.", "Pass"],
            ["IV-4", "Misspelled city", "Bot suggests checking the spelling.", "Pass"],
        ],
        "Input-validation test cases and results.",
        widths=[0.6, 2.3, 2.6, 0.7],
    )
    h2("15.5 Error-Handling Testing")
    table(
        ["ID", "Test case", "Expected result", "Status"],
        [
            ["EH-1", "Network failure during fetch", "Friendly network-error message shown.", "Pass"],
            ["EH-2", "Location permission denied", "Explanatory message; suggests typing a city.", "Pass"],
            ["EH-3", "Microphone permission denied", "Explanatory message shown.", "Pass"],
            ["EH-4", "Speech unsupported", "Feature degrades with a clear notice.", "Pass"],
        ],
        "Error-handling test cases and results.",
        widths=[0.6, 2.3, 2.6, 0.7],
    )
    h2("15.6 Browser-Compatibility Testing")
    table(
        ["ID", "Environment", "Outcome", "Status"],
        [
            ["BC-1", "Chrome (desktop)", "All features, including voice, work.", "Pass"],
            ["BC-2", "Edge (desktop)", "All features work.", "Pass"],
            ["BC-3", "Firefox (desktop)", "Core features work; speech limited.", "Pass*"],
            ["BC-4", "Safari (desktop/iOS)", "Core features work; speech limited.", "Pass*"],
            ["BC-5", "Mobile Chrome (Android)", "Responsive layout and GPS work.", "Pass"],
        ],
        "Browser-compatibility results (* feature-dependent, degrades gracefully).",
        widths=[0.6, 2.0, 2.9, 0.7],
    )


def ch_results():
    h1("16. Results and Discussion")
    body("The completed system meets all stated objectives. WeatherBot "
         "successfully interprets natural-language questions, accepts voice, GPS "
         "and map inputs, and returns current, next-day and seven-day forecasts "
         "for any geocodable location. All four input modes were verified to "
         "converge on the single report() routine, confirming the intended "
         "uniform internal contract.")
    body("Testing across functional, interface, integration, validation, error-"
         "handling and compatibility dimensions returned successful outcomes, "
         "with advanced voice features degrading gracefully where browser support "
         "is limited. The application performs its local processing "
         "instantaneously, with perceptible delay limited to the two asynchronous "
         "network calls, which are masked by a typing indicator.")
    body("The serverless, single-file architecture proved highly advantageous: "
         "the system was deployed to public static hosting with no configuration "
         "and operates continuously at the deployment URL at zero hosting cost. "
         "The discussion confirms that a polished, multi-modal assistant "
         "experience can indeed be achieved using only standard browser "
         "technologies and free public APIs.")


def ch_advantages():
    h1("17. Advantages")
    bullet(" Four interchangeable input modes (text, voice, GPS, map) within one "
           "consistent interface.", bold_lead="Multi-modal. ")
    bullet(" No back-end, database, build step or API keys; a single HTML file "
           "is the entire product.", bold_lead="Zero-configuration. ")
    bullet(" Deployable to any static host or opened directly from disk; "
           "trivially shareable.", bold_lead="Highly portable. ")
    bullet(" Local parsing and rendering are instantaneous; only data retrieval "
           "uses the network.", bold_lead="Fast. ")
    bullet(" No server to breach and no stored user data; explicit, in-browser "
           "consent for GPS and microphone.", bold_lead="Private and secure. ")
    bullet(" A clean, uniform internal contract makes new input modes or "
           "features easy to add.", bold_lead="Maintainable. ")
    bullet(" Free static hosting provides continuous 24/7 availability at no "
           "cost.", bold_lead="Always available. ")


def ch_limitations():
    h1("18. Limitations")
    bullet("The rule-based parser can misread very unusual phrasings or "
           "multi-word edge cases.")
    bullet("Only one place is matched per question (the top geocoding result).")
    bullet("Parsing vocabulary and the voice locale are English-only.")
    bullet("An internet connection is required for all live data; there is no "
           "offline cache.")
    bullet("Voice features depend on browser support and may be limited on some "
           "browsers.")
    bullet("The application is intentionally stateless, so no history or "
           "personalisation is retained between sessions.")


def ch_future():
    h1("19. Future Enhancements")
    bullet("Add multi-language support for both parsing and speech.")
    bullet("Offer hourly forecasts and severe-weather alerts.")
    bullet("Cache recent results for brief offline availability.")
    bullet("Introduce unit selection (°C/°F, km/h versus mph).")
    bullet("Allow comparison of multiple locations in a single query.")
    bullet("Optionally integrate a true language model for richer conversation.")
    bullet("Package the application as an installable Progressive Web App (PWA).")


def ch_conclusion():
    h1("20. Conclusion")
    body("WeatherBot demonstrates that a refined, multi-modal assistant can be "
         "built with nothing more than standard browser technologies and a "
         "handful of free public APIs. Its strength lies in a simple, uniform "
         "internal contract: every interaction — typed, spoken, located or "
         "tapped — is reduced to a location and a time intent, then handed to "
         "a single reporting routine that fetches live data and renders a "
         "friendly reply.")
    body("This clean separation of concerns makes the project easy to "
         "understand, easy to extend and effortless to deploy, while still "
         "delivering a polished, modern user experience. The system has been "
         "fully implemented, tested and publicly deployed, and is continuously "
         f"available at {URL}. WeatherBot thus fulfils its objectives and stands "
         "as a compact yet complete example of disciplined, end-to-end software "
         "engineering.")


def ch_references():
    h1("21. References")
    refs = [
        "Open-Meteo. Free Weather API and Geocoding API. https://open-meteo.com/",
        "BigDataCloud. Client-side Reverse Geocoding API. https://www.bigdatacloud.com/",
        "OpenStreetMap Contributors. OpenStreetMap. https://www.openstreetmap.org/",
        "Leaflet. An open-source JavaScript library for interactive maps. https://leafletjs.com/",
        "Mozilla Developer Network (MDN). Fetch API. https://developer.mozilla.org/docs/Web/API/Fetch_API",
        "Mozilla Developer Network (MDN). Geolocation API. https://developer.mozilla.org/docs/Web/API/Geolocation_API",
        "Mozilla Developer Network (MDN). Web Speech API. https://developer.mozilla.org/docs/Web/API/Web_Speech_API",
        "World Wide Web Consortium (W3C). HTML Living Standard and CSS Specifications. https://www.w3.org/",
        "GitHub. GitHub Pages static hosting. https://pages.github.com/",
        f"WeatherBot. Live deployment. {URL}",
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        run = p.add_run(f"[{i}] "); run.bold = True
        p.add_run(r)


# ---------------- footer with page numbers -----------------------------
def add_page_numbers():
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Page ")
    field(p, " PAGE ", "1")
    p.add_run(" of ")
    field(p, " NUMPAGES ", "1")
    for r in p.runs:
        r.font.size = Pt(9); r.font.color.rgb = GREY


# ======================== BUILD =======================================
cover()
certificate()
acknowledgement()
abstract()
toc_pages()
ch_introduction()
ch_problem()
ch_objectives()
ch_scope()
ch_literature()
ch_analysis()
ch_techstack()
ch_architecture()
ch_design()
ch_implementation()
ch_workflow()
ch_components()
ch_ui()
ch_accessibility()
ch_testing()
ch_results()
ch_advantages()
ch_limitations()
ch_future()
ch_conclusion()
ch_references()
add_page_numbers()

doc.save(OUT)
print("SAVED", os.path.abspath(OUT))
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
